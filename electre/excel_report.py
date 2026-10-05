"""Книга Excel с проверкой расчёта: формулы повторяют все шаги метода.

Программа записывает исходные данные и свои результаты, а все промежуточные величины
Excel считает собственными формулами. Совпадение чисел подтверждает правильность расчёта.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .data import HARRINGTON, WORD_VALUES, Task
from .methods import ElectreResult, Normalized, ParetoResult, Weights

TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
GOOD_FILL = PatternFill("solid", fgColor="E2EFDA")
CORE_FILL = PatternFill("solid", fgColor="FFF2CC")
WHITE_BOLD = Font(color="FFFFFF", bold=True, size=12)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="B7C7DA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _title(sheet, text: str, width: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(width, 2))
    cell = sheet.cell(row=1, column=1, value=text)
    cell.fill = TITLE_FILL
    cell.font = WHITE_BOLD
    cell.alignment = CENTER
    sheet.row_dimensions[1].height = 24


def _header_row(sheet, row: int, values: list[str], start: int = 1) -> None:
    for offset, value in enumerate(values):
        cell = sheet.cell(row=row, column=start + offset, value=value)
        cell.fill = HEAD_FILL
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = BORDER


def build_workbook(task: Task, normalized: Normalized, pareto_result: ParetoResult,
                   weights: Weights, result: ElectreResult) -> Workbook:
    """Сборка книги из шести листов."""
    book = Workbook()
    _sheet_source(book.active, task)
    _sheet_normalized(book.create_sheet("Приведение"), task, normalized)
    _sheet_weights(book.create_sheet("Веса"), task, weights)
    _sheet_pareto(book.create_sheet("Парето"), task, pareto_result)
    _sheet_electre(book.create_sheet("ЭЛЕКТРА"), task, result)
    _sheet_summary(book.create_sheet("Свод"), task, normalized, weights, result)
    return book


def _sheet_source(sheet, task: Task) -> None:
    """Исходные данные, словарь словесных оценок и баллы экспертов."""
    sheet.title = "Исходные данные"
    _title(sheet, task.title, len(task.alternatives) + 1)
    _header_row(sheet, 2, ["Критерий"] + task.alternatives)
    for index, (criterion, row) in enumerate(zip(task.criteria, task.values), start=3):
        name = criterion.name + (f", {criterion.unit}" if criterion.unit else "")
        sheet.cell(row=index, column=1, value=name).border = BORDER
        for column, value in enumerate(row, start=2):
            cell = sheet.cell(row=index, column=column, value=value)
            cell.alignment = CENTER
            cell.border = BORDER

    start = 3 + len(task.criteria) + 2
    sheet.cell(row=start - 1, column=1, value="Словесные оценки в числовом виде (шкала Харрингтона)").font = BOLD
    _header_row(sheet, start, ["Оценка", "Число"])
    for offset, (word, value) in enumerate(sorted(WORD_VALUES.items(), key=lambda item: item[1]), start=1):
        sheet.cell(row=start + offset, column=1, value=word).border = BORDER
        cell = sheet.cell(row=start + offset, column=2, value=value)
        cell.alignment = CENTER
        cell.border = BORDER
    words_end = start + len(WORD_VALUES)
    book = sheet.parent
    from openpyxl.workbook.defined_name import DefinedName
    book.defined_names.add(DefinedName("Словарь", attr_text=f"'Исходные данные'!$A${start + 1}:$B${words_end}"))

    scale_row = words_end + 2
    sheet.cell(row=scale_row, column=1, value="Диапазоны шкалы Харрингтона").font = BOLD
    _header_row(sheet, scale_row + 1, ["Качественная оценка", "От", "До"])
    for offset, (name, (low, high, _)) in enumerate(HARRINGTON.items(), start=1):
        sheet.cell(row=scale_row + 1 + offset, column=1, value=name).border = BORDER
        sheet.cell(row=scale_row + 1 + offset, column=2, value=low).border = BORDER
        sheet.cell(row=scale_row + 1 + offset, column=3, value=high).border = BORDER

    expert_row = scale_row + len(HARRINGTON) + 3
    sheet.cell(row=expert_row, column=1, value="Баллы экспертов (пятибалльная шкала)").font = BOLD
    _header_row(sheet, expert_row + 1, ["Эксперт"] + [c.name for c in task.criteria])
    for index, scores in enumerate(task.expert_scores, start=1):
        sheet.cell(row=expert_row + 1 + index, column=1, value=f"Эксперт {index}").border = BORDER
        for column, value in enumerate(scores, start=2):
            cell = sheet.cell(row=expert_row + 1 + index, column=column, value=value)
            cell.alignment = CENTER
            cell.border = BORDER
    for index, note in enumerate(task.expert_notes, start=1):
        sheet.cell(row=expert_row + 1 + len(task.expert_scores) + index, column=1, value=note)

    sheet.column_dimensions["A"].width = 46
    for column in range(2, len(task.alternatives) + 2):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.expert_row = expert_row  # запоминается для формул на других листах


def _sheet_normalized(sheet, task: Task, normalized: Normalized) -> None:
    """Перевод словесных оценок в числа и приведение к безразмерному виду."""
    _title(sheet, "Приведение оценок к безразмерному виду", len(task.alternatives) + 1)
    _header_row(sheet, 2, ["Критерий"] + task.alternatives)
    for index, criterion in enumerate(task.criteria, start=3):
        sheet.cell(row=index, column=1, value=criterion.name).border = BORDER
        for column in range(2, len(task.alternatives) + 2):
            letter = get_column_letter(column)
            if criterion.kind == "качественный":
                formula = f"=VLOOKUP('Исходные данные'!{letter}{index},Словарь,2,FALSE)"
            else:
                formula = f"='Исходные данные'!{letter}{index}"
            cell = sheet.cell(row=index, column=column, value=formula)
            cell.number_format = "0.00"
            cell.alignment = CENTER
            cell.border = BORDER

    base = 3 + len(task.criteria) + 1
    sheet.cell(row=base - 1, column=1, value="Безразмерные оценки").font = BOLD
    _header_row(sheet, base, ["Критерий"] + task.alternatives)
    for index, criterion in enumerate(task.criteria):
        row = base + 1 + index
        source = 3 + index
        first = get_column_letter(2)
        last = get_column_letter(len(task.alternatives) + 1)
        sheet.cell(row=row, column=1, value=f"{criterion.name} ({criterion.better} — лучше)").border = BORDER
        for column in range(2, len(task.alternatives) + 2):
            letter = get_column_letter(column)
            if criterion.maximize:
                formula = f"={letter}{source}/MAX(${first}${source}:${last}${source})"
            else:
                formula = f"=MIN(${first}${source}:${last}${source})/{letter}{source}"
            cell = sheet.cell(row=row, column=column, value=formula)
            cell.number_format = "0.000"
            cell.alignment = CENTER
            cell.border = BORDER
    sheet.column_dimensions["A"].width = 42
    for column in range(2, len(task.alternatives) + 2):
        sheet.column_dimensions[get_column_letter(column)].width = 14
    sheet.base_row = base + 1


def _sheet_weights(sheet, task: Task, weights: Weights) -> None:
    """Метод ранга: суммы баллов и веса критериев."""
    _title(sheet, "Веса критериев по методу ранга", 5)
    _header_row(sheet, 2, ["Критерий", "Эксперт 1", "Эксперт 2", "Сумма баллов", "Вес критерия"])
    expert_row = 3 + len(task.criteria) + 2 + len(WORD_VALUES) + 2 + len(HARRINGTON) + 3
    for index, criterion in enumerate(task.criteria, start=3):
        sheet.cell(row=index, column=1, value=criterion.name).border = BORDER
        for expert in range(len(task.expert_scores)):
            source_row = expert_row + 2 + expert
            source_column = get_column_letter(index - 1)
            cell = sheet.cell(row=index, column=2 + expert,
                              value=f"='Исходные данные'!{source_column}{source_row}")
            cell.alignment = CENTER
            cell.border = BORDER
        total_column = get_column_letter(1 + len(task.expert_scores))
        sheet.cell(row=index, column=4, value=f"=SUM(B{index}:{total_column}{index})").alignment = CENTER
        last = 2 + len(task.criteria)
        sheet.cell(row=index, column=5, value=f"=D{index}/SUM($D$3:$D${last})").number_format = "0.000"
        for column in range(1, 6):
            sheet.cell(row=index, column=column).border = BORDER

    last = 2 + len(task.criteria)
    sheet.cell(row=last + 1, column=1, value="Итого").font = BOLD
    sheet.cell(row=last + 1, column=4, value=f"=SUM(D3:D{last})").font = BOLD
    sheet.cell(row=last + 1, column=5, value=f"=SUM(E3:E{last})").font = BOLD
    sheet.cell(row=last + 3, column=1,
               value="Вес критерия — это доля его баллов в общей сумме баллов, выставленных экспертами.")
    sheet.column_dimensions["A"].width = 34
    for column in "BCDE":
        sheet.column_dimensions[column].width = 16


def _sheet_pareto(sheet, task: Task, pareto_result: ParetoResult) -> None:
    """Проверка множества Парето формулами."""
    count = len(task.alternatives)
    _title(sheet, "Множество Парето: какая площадка кого превосходит", count + 2)
    sheet.cell(row=2, column=1, value="Строка превосходит столбец («1» — превосходит по всем критериям)").font = BOLD
    _header_row(sheet, 3, [""] + task.alternatives)
    criteria_count = len(task.criteria)
    first_row = 3 + criteria_count + 2  # первая строка блока безразмерных оценок на листе «Приведение»
    for i in range(count):
        row = 4 + i
        sheet.cell(row=row, column=1, value=task.alternatives[i]).font = BOLD
        sheet.cell(row=row, column=1).border = BORDER
        source_i = get_column_letter(2 + i)
        for j in range(count):
            column = 2 + j
            source_j = get_column_letter(2 + j)
            if i == j:
                cell = sheet.cell(row=row, column=column, value="—")
            else:
                formula = (
                    f"=IF(AND(SUMPRODUCT(--(Приведение!${source_i}${first_row}:${source_i}${first_row + criteria_count - 1}"
                    f">=Приведение!${source_j}${first_row}:${source_j}${first_row + criteria_count - 1}))={criteria_count},"
                    f"SUMPRODUCT(--(Приведение!${source_i}${first_row}:${source_i}${first_row + criteria_count - 1}"
                    f">Приведение!${source_j}${first_row}:${source_j}${first_row + criteria_count - 1}))>0),1,0)"
                )
                cell = sheet.cell(row=row, column=column, value=formula)
            cell.alignment = CENTER
            cell.border = BORDER

    result_row = 4 + count + 1
    sheet.cell(row=result_row, column=1, value="Входит в множество Парето").font = BOLD
    for j in range(count):
        column = get_column_letter(2 + j)
        cell = sheet.cell(row=result_row, column=2 + j,
                          value=f'=IF(SUM({column}4:{column}{3 + count})=0,"да","нет")')
        cell.alignment = CENTER
        cell.border = BORDER
        cell.fill = GOOD_FILL
    sheet.cell(row=result_row + 1, column=1, value="Результат программы").font = BOLD
    for j in range(count):
        cell = sheet.cell(row=result_row + 1, column=2 + j, value="да" if j in pareto_result.kept else "нет")
        cell.alignment = CENTER
        cell.border = BORDER
    check = sheet.cell(
        row=result_row + 3, column=1,
        value=f'=IF(SUMPRODUCT(--(B{result_row}:{get_column_letter(1 + count)}{result_row}='
              f'B{result_row + 1}:{get_column_letter(1 + count)}{result_row + 1}))={count},'
              f'"Совпадает: Excel отобрал те же площадки","Есть расхождения")')
    check.font = BOLD
    check.fill = GOOD_FILL
    sheet.column_dimensions["A"].width = 32
    for column in range(2, count + 2):
        sheet.column_dimensions[get_column_letter(column)].width = 12


def _sheet_electre(sheet, task: Task, result: ElectreResult) -> None:
    """Матрицы согласия и несогласия, посчитанные формулами Excel."""
    size = len(result.names)
    _title(sheet, "Метод ЭЛЕКТРА: индексы согласия и несогласия", size + 2)
    criteria_count = len(task.criteria)
    first_row = 3 + criteria_count + 2  # блок безразмерных оценок на листе «Приведение»
    columns = {a: get_column_letter(2 + result.indices[a]) for a in range(size)}

    sheet.cell(row=2, column=1, value="Индексы согласия: сумма весов критериев, по которым строка не хуже столбца").font = BOLD
    _header_row(sheet, 3, [""] + result.names)
    for a in range(size):
        row = 4 + a
        sheet.cell(row=row, column=1, value=result.names[a]).font = BOLD
        sheet.cell(row=row, column=1).border = BORDER
        for b in range(size):
            cell = sheet.cell(row=row, column=2 + b)
            if a != b:
                rng_a = f"Приведение!${columns[a]}${first_row}:${columns[a]}${first_row + criteria_count - 1}"
                rng_b = f"Приведение!${columns[b]}${first_row}:${columns[b]}${first_row + criteria_count - 1}"
                cell.value = f"=SUMPRODUCT(--({rng_a}>={rng_b}),Веса!$E$3:$E${2 + criteria_count})"
                cell.number_format = "0.000"
            else:
                cell.value = "—"
            cell.alignment = CENTER
            cell.border = BORDER

    start = 4 + size + 2
    sheet.cell(row=start - 1, column=1, value="Индексы несогласия: наибольшее отставание строки от столбца").font = BOLD
    _header_row(sheet, start, [""] + result.names)
    for a in range(size):
        row = start + 1 + a
        sheet.cell(row=row, column=1, value=result.names[a]).font = BOLD
        sheet.cell(row=row, column=1).border = BORDER
        for b in range(size):
            cell = sheet.cell(row=row, column=2 + b)
            if a != b:
                rng_a = f"Приведение!${columns[a]}${first_row}:${columns[a]}${first_row + criteria_count - 1}"
                rng_b = f"Приведение!${columns[b]}${first_row}:${columns[b]}${first_row + criteria_count - 1}"
                cell.value = f"=SUMPRODUCT(MAX(0,{rng_b}-{rng_a}))"
                cell.number_format = "0.000"
            else:
                cell.value = "—"
            cell.alignment = CENTER
            cell.border = BORDER
    sheet.column_dimensions["A"].width = 28
    for column in range(2, size + 2):
        sheet.column_dimensions[get_column_letter(column)].width = 12
    sheet.concordance_row = 4
    sheet.discordance_row = start + 1


def _sheet_summary(sheet, task: Task, normalized: Normalized, weights: Weights, result: ElectreResult) -> None:
    """Сравнение результатов программы и формул Excel, пороги и «ядро»."""
    size = len(result.names)
    _title(sheet, "Итог: сравнение расчёта программы и формул Excel", 7)
    _header_row(sheet, 2, [
        "Площадка",
        "Мин. индекс согласия (Excel)",
        "Мин. индекс согласия (программа)",
        "Макс. индекс несогласия (Excel)",
        "Макс. индекс несогласия (программа)",
        "Вывод Excel",
        "Вывод программы",
    ])
    conc_row, disc_row = 4, 4 + size + 3
    for a in range(size):
        row = 3 + a
        letters = get_column_letter(2 + size)
        sheet.cell(row=row, column=1, value=result.names[a]).font = BOLD
        sheet.cell(row=row, column=2, value=f"=MIN(ЭЛЕКТРА!B{conc_row + a}:{letters}{conc_row + a})").number_format = "0.000"
        sheet.cell(row=row, column=3, value=round(result.min_concordance[a], 6)).number_format = "0.000"
        sheet.cell(row=row, column=4, value=f"=MAX(ЭЛЕКТРА!B{disc_row + a}:{letters}{disc_row + a})").number_format = "0.000"
        sheet.cell(row=row, column=5, value=round(result.max_discordance[a], 6)).number_format = "0.000"
        sheet.cell(row=row, column=6,
                   value=f'=IF(AND(B{row}>=$B$%d,D{row}<=$B$%d),"входит в ядро","—")' % (size + 5, size + 6))
        sheet.cell(row=row, column=7, value="входит в ядро" if a in result.core else "—")
        for column in range(1, 8):
            sheet.cell(row=row, column=column).alignment = CENTER
            sheet.cell(row=row, column=column).border = BORDER
        if a in result.core:
            for column in range(1, 8):
                sheet.cell(row=row, column=column).fill = CORE_FILL

    threshold_c_row = size + 5
    threshold_d_row = size + 6
    sheet.cell(row=threshold_c_row, column=1, value="Порог индекса согласия c*").font = BOLD
    sheet.cell(row=threshold_c_row, column=2, value=result.threshold_c).font = BOLD
    sheet.cell(row=threshold_d_row, column=1, value="Порог индекса несогласия d*").font = BOLD
    sheet.cell(row=threshold_d_row, column=2, value=result.threshold_d).font = BOLD
    sheet.cell(row=threshold_d_row + 1, column=1,
               value="Пороги можно менять прямо здесь: столбец «Вывод Excel» пересчитается сам.")

    last = 2 + size
    check_row = threshold_d_row + 3
    sheet.cell(row=check_row, column=1, value="Проверка совпадения").font = BOLD
    cell = sheet.cell(
        row=check_row, column=2,
        value=f'=IF(AND(SUMPRODUCT(--(ABS(B3:B{last}-C3:C{last})>0.001))=0,'
              f'SUMPRODUCT(--(ABS(D3:D{last}-E3:E{last})>0.001))=0),'
              f'"Совпадает: формулы Excel дают те же индексы","Есть расхождения")')
    cell.font = BOLD
    cell.fill = GOOD_FILL
    sheet.cell(row=check_row + 1, column=1, value="Ядро по расчёту программы").font = BOLD
    sheet.cell(row=check_row + 1, column=2, value=", ".join(result.core_names()) or "пусто").font = BOLD
    sheet.cell(row=check_row + 1, column=2).fill = CORE_FILL

    sheet.cell(row=check_row + 3, column=1, value="Веса критериев (метод ранга)").font = BOLD
    for index, (criterion, weight) in enumerate(zip(task.criteria, weights.values), start=1):
        sheet.cell(row=check_row + 3 + index, column=1, value=criterion.name)
        sheet.cell(row=check_row + 3 + index, column=2, value=round(weight, 6)).number_format = "0.000"
        sheet.cell(row=check_row + 3 + index, column=3, value=f"=Веса!E{2 + index}").number_format = "0.000"
    sheet.cell(row=check_row + 3, column=3, value="значение из формул Excel").font = BOLD

    sheet.column_dimensions["A"].width = 34
    for column in "BCDEFG":
        sheet.column_dimensions[column].width = 22


def save_workbook(task: Task, normalized: Normalized, pareto_result: ParetoResult,
                  weights: Weights, result: ElectreResult, path: str) -> str:
    """Сохранение книги на диск."""
    book = build_workbook(task, normalized, pareto_result, weights, result)
    book.save(path)
    return path
